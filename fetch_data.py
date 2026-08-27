from __future__ import annotations

import json
import re
import ssl
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
from dataclasses import dataclass
from datetime import datetime, timezone
from html import unescape
from pathlib import Path
from typing import Any
from urllib.request import Request, urlopen

from openpyxl import Workbook
from openpyxl.styles import Font

# EU 2026/1457 Annex I - flat steel product categories (1.A - 10).
# Order numbers verified against the regulation's Annex I table.
PRODUCT_CATEGORIES = [
    {
        "key": "1A",
        "quota_section": "1.A",
        "product_group": "Non Alloy and Other Alloy Hot Rolled Sheets and Strips",
        "exporters": {
            "099801": "Türkiye",
            "099802": "Japan",
            "099803": "India",
            "099804": "Taiwan",
            "099805": "Ukraine",
            "099806": "Korea",
            "099807": "Viet Nam",
            "099808": "Egypt",
            "099809": "Serbia",
            "099500": "FTA Quota - CSQ",
            "099701": "Brazil",
            "099705": "United Kingdom",
            "099702": "Indonesia",
            "099810": "Australia",
            "099811": "Saudi Arabia",
            "099704": "Switzerland",
            "099812": "Kazakhstan",
            "099703": "North Macedonia",
            "099600": "Other countries",
            "099700": "FTA Quota - Other countries",
        },
    },
    {
        "key": "1B",
        "quota_section": "1.B",
        "product_group": "Non Alloy and Other Alloy Hot Rolled Sheets and Strips",
        "exporters": {
            "099813": "United Kingdom",
            "099814": "United States",
            "099815": "Japan",
            "099816": "China",
            "099501": "FTA Quota - CSQ",
            "099601": "Other countries",
            "099706": "FTA Quota - Other countries",
        },
    },
    {
        "key": "2",
        "quota_section": "2",
        "product_group": "Non Alloy and Other Alloy Cold Rolled Sheets",
        "exporters": {
            "099817": "Taiwan",
            "099818": "India",
            "099819": "Korea",
            "099820": "Türkiye",
            "099821": "United Kingdom",
            "099822": "Japan",
            "099823": "Ukraine",
            "099502": "FTA Quota - CSQ",
            "099602": "Other countries",
            "099707": "FTA Quota - Other countries",
            "099709": "Egypt",
            "099710": "Switzerland",
            "099708": "Brazil",
        },
    },
    {
        "key": "3A",
        "quota_section": "3.A",
        "product_group": "Electrical Sheets (other than GOES)",
        "exporters": {
            "099824": "Japan",
            "099825": "United Kingdom",
            "099826": "China",
            "099827": "Türkiye",
            "099503": "FTA Quota - CSQ",
            "099603": "Other countries",
            "099711": "FTA Quota - Other countries",
        },
    },
    {
        "key": "3B",
        "quota_section": "3.B",
        "product_group": "Electrical Sheets (other than GOES)",
        "exporters": {
            "099828": "China",
            "099829": "Taiwan",
            "099830": "Korea",
            "099831": "Viet Nam",
            "099504": "FTA Quota - CSQ",
            "099604": "Other countries",
            "099712": "FTA Quota - Other countries",
            "099713": "Japan",
        },
    },
    {
        "key": "4A",
        "quota_section": "4.A",
        "product_group": "Metallic Coated Sheets",
        "exporters": {
            "099832": "Viet Nam",
            "099833": "Taiwan",
            "099834": "Turkiye",
            "099835": "India",
            "099836": "Korea",
            "099505": "FTA Quota - CSQ",
            "099605": "Other countries",
            "099714": "FTA Quota - Other countries",
            "099718": "United Kingdom",
            "099716": "Japan",
            "099715": "Egypt",
            "099717": "South Africa",
        },
    },
    {
        "key": "4B",
        "quota_section": "4.B",
        "product_group": "Metallic Coated Sheets",
        "exporters": {
            "099837": "Korea",
            "099838": "China",
            "099839": "United Kingdom",
            "099840": "Türkiye",
            "099841": "India",
            "099506": "FTA Quota - CSQ",
            "099606": "Other countries",
            "099719": "FTA Quota - Other countries",
            "099720": "Egypt",
            "099721": "Switzerland",
        },
    },
    {
        "key": "5",
        "quota_section": "5",
        "product_group": "Organic Coated Sheets",
        "exporters": {
            "099842": "India",
            "099843": "Korea",
            "099844": "Viet Nam",
            "099845": "Türkiye",
            "099846": "Taiwan",
            "099847": "United Kingdom",
            "099507": "FTA Quota - CSQ",
            "099607": "Other countries",
            "099722": "FTA Quota - Other countries",
            "099723": "North Macedonia",
        },
    },
    {
        "key": "6",
        "quota_section": "6",
        "product_group": "Tin Mill Products",
        "exporters": {
            "099848": "China",
            "099849": "Serbia",
            "099850": "United Kingdom",
            "099851": "Türkiye",
            "099852": "Korea",
            "099853": "India",
            "099508": "FTA Quota - CSQ",
            "099608": "Other countries",
            "099724": "FTA Quota - Other countries",
            "099725": "Japan",
            "099726": "Singapore",
        },
    },
    {
        "key": "7",
        "quota_section": "7",
        "product_group": "Non Alloy and Other Alloy Quarto Plates",
        "exporters": {
            "099854": "Korea",
            "099855": "Indonesia",
            "099856": "India",
            "099857": "Japan",
            "099858": "North Macedonia",
            "099859": "United Kingdom",
            "099509": "FTA Quota - CSQ",
            "099609": "Other countries",
            "099727": "FTA Quota - Other countries",
            "099728": "Türkiye",
            "099491": "United Kingdom (Northern Ireland)",
        },
    },
    {
        "key": "8",
        "quota_section": "8",
        "product_group": "Stainless Hot Rolled Sheets and Strips",
        "exporters": {
            "099860": "Taiwan",
            "099861": "Indonesia",
            "099862": "India",
            "099863": "China",
            "099864": "Korea",
            "099865": "Türkiye",
            "099866": "South Africa",
            "099510": "FTA Quota - CSQ",
            "099610": "Other countries",
            "099729": "FTA Quota - Other countries",
            "099492": "United Kingdom (Northern Ireland)",
        },
    },
    {
        "key": "9",
        "quota_section": "9",
        "product_group": "Stainless Cold Rolled Sheets and Strips",
        "exporters": {
            "099867": "Taiwan",
            "099868": "China",
            "099869": "Korea",
            "099870": "Türkiye",
            "099871": "South Africa",
            "099872": "Viet Nam",
            "099873": "India",
            "099511": "FTA Quota - CSQ",
            "099611": "Other countries",
            "099730": "FTA Quota - Other countries",
            "099731": "Switzerland",
            "099493": "United Kingdom (Northern Ireland)",
        },
    },
    {
        "key": "10",
        "quota_section": "10",
        "product_group": "Stainless Hot Rolled Quarto Plates",
        "exporters": {
            "099874": "China",
            "099875": "India",
            "099876": "Korea",
            "099877": "South Africa",
            "099512": "FTA Quota - CSQ",
            "099612": "Other countries",
            "099732": "FTA Quota - Other countries",
        },
    },
]

BASE_URL = "https://ec.europa.eu/taxation_customs/dds2/taric"
LIST_URL = (
    BASE_URL
    + "/quota_list.jsp?Lang=en&Code={code}&Year=2026&Expand=true&Offset=0"
)
DETAIL_URL = BASE_URL + "/quota_tariff_details.jsp?Lang=en&StartDate={start_date}&Code={code}"
DATA_DIR = Path(__file__).resolve().parent / "public" / "data"
OUT_PATH = DATA_DIR / "orders.json"
XLSX_PATH = DATA_DIR / "eu-steel-trq-flat-dashboard.xlsx"


@dataclass
class Quantity:
    value: float | None
    unit: str | None
    raw: str


def fetch_text(url: str) -> str:
    request = Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0",
            "Accept-Language": "en-US,en;q=0.9",
        },
    )
    ssl_context = ssl._create_unverified_context()
    with urlopen(request, timeout=30, context=ssl_context) as response:
        raw = response.read()
        charset = response.headers.get_content_charset() or "utf-8"
        return raw.decode(charset, errors="replace")


def squash_whitespace(text: str) -> str:
    text = text.replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def clean_html_text(fragment: str) -> str:
    fragment = re.sub(r"<br\s*/?>", "\n", fragment, flags=re.IGNORECASE)
    fragment = re.sub(r"</div\s*>", "\n", fragment, flags=re.IGNORECASE)
    fragment = re.sub(r"<div[^>]*>", "", fragment, flags=re.IGNORECASE)
    fragment = re.sub(r"</?a[^>]*>", "", fragment, flags=re.IGNORECASE)
    fragment = re.sub(r"<[^>]+>", " ", fragment)
    lines = [squash_whitespace(unescape(line)) for line in fragment.splitlines()]
    return "\n".join(line for line in lines if line)


def parse_quantity(text: str) -> Quantity:
    raw = squash_whitespace(text)
    match = re.match(r"^([0-9]+(?:[.,][0-9]+)?)\s+(.+)$", raw)
    if not match:
        number_match = re.match(r"^([0-9]+(?:[.,][0-9]+)?)$", raw)
        if number_match:
            return Quantity(float(number_match.group(1).replace(",", "")), None, raw)
        return Quantity(None, None, raw)
    value = float(match.group(1).replace(",", ""))
    unit = match.group(2).strip()
    return Quantity(value, unit, raw)


def find_first(pattern: str, text: str) -> str | None:
    match = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
    return match.group(1) if match else None


def parse_table_rows(container_html: str) -> dict[str, str]:
    rows: dict[str, str] = {}
    row_pattern = re.compile(
        r"<tr[^>]*class=\"ecl-table__row\"[^>]*>(.*?)</tr>",
        flags=re.IGNORECASE | re.DOTALL,
    )
    cell_pattern = re.compile(r"<td[^>]*>(.*?)</td>", flags=re.IGNORECASE | re.DOTALL)
    for row_html in row_pattern.findall(container_html):
        cells = cell_pattern.findall(row_html)
        if len(cells) < 2:
            continue
        label_html, value_html = cells[0], cells[1]
        label = clean_html_text(label_html)
        label = re.sub(r"\(\s*indicative\s*\)", "", label, flags=re.IGNORECASE)
        label = squash_whitespace(label)
        value = clean_html_text(value_html)
        rows[label] = value
    return rows


def parse_list_page(code: str, html: str) -> dict[str, Any]:
    tbody_match = re.search(
        r"<tbody[^>]*class=\"ecl-table__body\"[^>]*>(.*?)</tbody>",
        html,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if not tbody_match:
        raise RuntimeError(f"No list body found for order number {code}")
    row_match = re.search(
        r"<tr[^>]*class=\"ecl-table__row\"[^>]*>(.*?)</tr>",
        tbody_match.group(1),
        flags=re.IGNORECASE | re.DOTALL,
    )
    if not row_match:
        raise RuntimeError(f"No list row found for order number {code}")
    row_html = row_match.group(1)

    cells = re.findall(r"<td[^>]*>(.*?)</td>", row_html, flags=re.IGNORECASE | re.DOTALL)
    if len(cells) < 6:
        raise RuntimeError(f"Unexpected list row structure for order number {code}")

    detail_href = find_first(r'href="([^"]*quota_tariff_details\.jsp[^"]+)"', row_html)
    start_date = find_first(r"StartDate=([0-9]{4}-[0-9]{2}-[0-9]{2})", detail_href or "")
    last_update = find_first(r"Last TARIC update:&nbsp;([0-9]{2}-[0-9]{2}-[0-9]{4})", html)

    return {
        "order_number": clean_html_text(cells[0]),
        "origins_summary": clean_html_text(cells[1]),
        "start_date": clean_html_text(cells[2]),
        "end_date": clean_html_text(cells[3]),
        "balance": parse_quantity(clean_html_text(cells[4])).__dict__,
        "detail_url": detail_href,
        "detail_start_date": start_date,
        "source_last_taric_update": last_update,
    }


def parse_detail_page(code: str, html: str) -> dict[str, Any]:
    start = html.find('<div id="quotaDetailsMarkedUpContainer"')
    marker = html.find("Associated TARIC code", start)
    end = html.find("</table>", marker)
    container = html[start:end] if start >= 0 and marker >= 0 and end >= 0 else None
    if not container:
        raise RuntimeError(f"No detail container found for order number {code}")

    rows = parse_table_rows(container)
    last_update = find_first(r"Last TARIC update:&nbsp;([0-9]{2}-[0-9]{2}-[0-9]{4})", html)

    initial_amount = parse_quantity(rows.get("Initial amount", ""))
    current_amount = parse_quantity(rows.get("Amount", ""))
    balance = parse_quantity(rows.get("Balance", ""))

    used_value = None
    if initial_amount.value is not None and balance.value is not None:
        used_value = round(initial_amount.value - balance.value, 2)

    utilization_pct = None
    if initial_amount.value and used_value is not None:
        utilization_pct = round((used_value / initial_amount.value) * 100, 2)

    return {
        "order_number": rows.get("Order number", code),
        "validity_period": rows.get("Validity period", ""),
        "origin": rows.get("Origin", ""),
        "initial_amount": initial_amount.__dict__,
        "current_amount": current_amount.__dict__,
        "balance": balance.__dict__,
        "used_amount": {
            "value": used_value,
            "unit": initial_amount.unit or balance.unit,
            "raw": "" if used_value is None else f"{used_value}",
        },
        "utilization_pct": utilization_pct,
        "exhaustion_date": rows.get("Exhaustion date", ""),
        "critical": rows.get("Critical", ""),
        "last_import_date": rows.get("Last import date", ""),
        "last_allocation_date": rows.get("Last allocation date", ""),
        "awaiting_allocation": rows.get("Total awaiting allocation", ""),
        "awaiting_allocation_mt": to_mt(
            parse_quantity(rows.get("Total awaiting allocation", "")).__dict__
        ),
        "blocking_period": rows.get("Blocking period", ""),
        "suspension_period": rows.get("Suspension period", ""),
        "allocated_pct_last_allocation": rows.get("Allocated percentage at the last allocation", ""),
        "associated_taric_codes": rows.get("Associated TARIC code", "").splitlines(),
        "source_last_taric_update": last_update,
    }


def to_mt(quantity: dict[str, Any]) -> dict[str, Any]:
    value = quantity.get("value")
    if value is None:
        return {"value": None, "unit": "MT", "raw": ""}
    converted = round(value / 1000, 2)
    return {"value": converted, "unit": "MT", "raw": f"{converted:.2f} MT"}


def build_record(code: str, quota_section: str, product_group: str, exporter: str) -> dict[str, Any]:
    list_html = fetch_text(LIST_URL.format(code=code))
    list_data = parse_list_page(code, list_html)
    if not list_data["detail_start_date"]:
        raise RuntimeError(f"Missing detail start date for order number {code}")

    detail_html = fetch_text(
        DETAIL_URL.format(code=code, start_date=list_data["detail_start_date"])
    )
    detail_data = parse_detail_page(code, detail_html)

    source_last_update = detail_data["source_last_taric_update"] or list_data["source_last_taric_update"]
    return {
        "order_number": code,
        "quota_section": quota_section,
        "product_group": product_group,
        "exporter": exporter,
        "origins_summary": list_data["origins_summary"],
        "start_date": list_data["start_date"],
        "end_date": list_data["end_date"],
        "detail_url": list_data["detail_url"],
        "source_last_taric_update": source_last_update,
        "initial_amount_mt": to_mt(detail_data["initial_amount"]),
        "used_amount_mt": to_mt(detail_data["used_amount"]),
        "balance_mt": to_mt(detail_data["balance"]),
        **detail_data,
    }


def build_category(category: dict[str, Any]) -> dict[str, Any]:
    records = []
    for code, exporter in category["exporters"].items():
        try:
            records.append(
                build_record(code, category["quota_section"], category["product_group"], exporter)
            )
        except Exception as exc:  # keep going even if a single order number fails
            print(f"  ! Skipped {category['key']}/{code} ({exporter}): {exc}")

    report_period = ""
    if records:
        report_period = f'{records[0]["start_date"]} - {records[0]["end_date"]}'

    return {
        "key": category["key"],
        "quota_section": category["quota_section"],
        "product_group": category["product_group"],
        "order_numbers": list(category["exporters"].keys()),
        "report_period": report_period,
        "items": records,
    }


def build_excel(categories_payload: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    headers = [
        "Quota Section",
        "Product Group",
        "Order Number",
        "Exporter",
        "Start Date",
        "End Date",
        "Initial (MT)",
        "Used (MT)",
        "Balance (MT)",
        "Utilization (%)",
        "Critical",
        "Last Allocation Date",
        "Awaiting Allocation (MT)",
        "Detail URL",
    ]

    for category in categories_payload:
        sheet_title = f"{category['key']} {category['product_group']}"[:31]
        sheet = workbook.create_sheet(title=sheet_title)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)

        for item in category["items"]:
            sheet.append(
                [
                    item["quota_section"],
                    item["product_group"],
                    item["order_number"],
                    item["exporter"],
                    item["start_date"],
                    item["end_date"],
                    item["initial_amount_mt"]["value"],
                    item["used_amount_mt"]["value"],
                    item["balance_mt"]["value"],
                    item["utilization_pct"],
                    item["critical"],
                    item["last_allocation_date"],
                    item["awaiting_allocation_mt"]["value"],
                    item["detail_url"],
                ]
            )

        for row in sheet.iter_rows(min_row=2, min_col=7, max_col=10):
            for cell in row:
                cell.number_format = "0.00"

        widths = {
            "A": 14, "B": 34, "C": 14, "D": 24, "E": 14, "F": 14,
            "G": 14, "H": 14, "I": 14, "J": 14, "K": 12, "L": 20,
            "M": 18, "N": 80,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    workbook.save(XLSX_PATH)


def main() -> None:
    categories_payload = []
    for category in PRODUCT_CATEGORIES:
        print(f"Fetching category {category['key']} ({category['product_group']})...")
        categories_payload.append(build_category(category))

    payload = {
        "title": "EU STEEL TRQ 판재류 소진 현황",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "source": "https://ec.europa.eu/taxation_customs/dds2/taric/quota_consultation.jsp?Lang=en",
        "year_filter": 2026,
        "categories": categories_payload,
    }

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    build_excel(categories_payload)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
